# -*- coding: utf-8 -*-
#!/usr/bin/env python3
"""
Módulo Especialista: Coupa Enterprise (Trabalhador do Maestro)
Autor: Allan Simão
"""
import os
import re
import time
import json
import logging
import shutil
import tempfile
import sys
import threading
import random
import traceback
from datetime import datetime
from typing import List, Dict, Set, Optional, Any
from google import genai
from groq import Groq

import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Side, PatternFill

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.options import Options
from selenium.common.exceptions import NoAlertPresentException, NoSuchElementException
from selenium.webdriver.common.action_chains import ActionChains

from docx import Document
from docx.shared import Pt, RGBColor, Cm

import base64

logger = logging.getLogger("CoupaApp")
logger.setLevel(logging.INFO)
formatter = logging.Formatter('%(asctime)s [%(levelname)s] %(message)s', datefmt='%H:%M:%S')

if not logger.handlers:
    console_handler = logging.StreamHandler(sys.stdout)
    console_handler.setFormatter(formatter)
    logger.addHandler(console_handler)


class UtilitariosMixin:
    def verificar_acesso_pastas(self) -> bool:
        logger.info(" [PASSO 1] Verificando estrutura de pastas...")
        pastas = [self.PASTA_RELATORIOS, self.PASTA_BLOQUEADOS, self.PASTA_DATABASE, self.PASTA_BACKUPS, self.PASTA_LOGS]
        for p in pastas:
            try:
                if not os.path.exists(p): os.makedirs(p, exist_ok=True)
            except Exception as e:
                logger.error(f" Erro ao criar pasta {p}: {e}")
                return False
        
        self._configurar_log_arquivo()

        try:
            if not os.path.exists(self.ARQUIVO_MARCAS_BLOQUEADAS):
                pd.DataFrame({"Marcas Bloqueadas": ["SCHNEIDER", "SIEMENS"]}).to_excel(self.ARQUIVO_MARCAS_BLOQUEADAS, index=False)
            if not os.path.exists(self.ARQUIVO_BASE_MARCAS):
                pd.DataFrame({"Marcas Conhecidas": ["WEG", "BOSCH"]}).to_excel(self.ARQUIVO_BASE_MARCAS, index=False)
            if not os.path.exists(self.ARQUIVO_PLANILHA_CONTROLE):
                wb = Workbook(); ws = wb.active; ws.title = self.NOME_ABA_ALVO
                ws.append(["COTAÇÃO", "VENCIMENTO", "ITEM", "QUANTIDADE", "LOCALIDADE", "VENDEDOR", "MODELOS", "MARCAS", "RESPOSTA"])
                wb.save(self.ARQUIVO_PLANILHA_CONTROLE)
            if not os.path.exists(self.ARQUIVO_JSONL):
                with open(self.ARQUIVO_JSONL, 'w', encoding='utf-8') as f: f.write("")
        except Exception as e:
            logger.error(f" Erro arquivos base: {e}")
            return False
        return True

    def _configurar_log_arquivo(self):
        try:
            # Caminho fixo exigido
            pasta_logs_fixa = r"\\SERVIDOR2\Publico\ALLAN\Logs"
            os.makedirs(pasta_logs_fixa, exist_ok=True)
            caminho_log = os.path.join(pasta_logs_fixa, "log-coupa.txt")
            
            # Remove arquivos de log antigos diários, se existirem na memória
            handlers_to_remove = [h for h in logger.handlers if isinstance(h, logging.FileHandler)]
            for h in handlers_to_remove:
                logger.removeHandler(h)

            # Define o modo 'a' (Append) para escrever continuamente no mesmo txt
            file_handler = logging.FileHandler(caminho_log, mode='a', encoding='utf-8')
            file_handler.setFormatter(formatter)
            logger.addHandler(file_handler)
            logger.info(f" Log contínuo do Coupa ativo em: {caminho_log}")
        except Exception as e:
            pass

    def _validar_txt(self, t):
        return not (re.search(r'\d{2}/\d{2}/\d{4}', t) or len(t) < 2)

    def parse_local(self, t):
        t = t.replace("Local de entrega:", "").strip()
        p = [x.strip() for x in t.split('-')]
        ufs = {"AC","AL","AP","AM","BA","CE","DF","ES","GO","MA","MT","MS","MG","PA","PB","PR","PE","PI","RJ","RN","RS","RO","RR","SC","SP","SE","TO"}
        for i in range(len(p)-1, 0, -1):
            if p[i].upper().replace(".", "") in ufs:
                c = p[i-1]
                if c.upper() not in ["S/N", "S/N."]: return f"{c} - {p[i].upper().replace('.', '')}"
        return t

    def _filter_desc(self, t):
        if "PT ||" in t:
            try:
                s = t.index("PT ||")
                ends = [x for x in [t.find("***", s), t.find("ES ||", s), t.find("EN ||", s)] if x != -1]
                return t[s:min(ends) if ends else len(t)].strip()
            except: pass
        return t
    
