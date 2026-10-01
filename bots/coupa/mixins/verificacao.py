# -*- coding: utf-8 -*-
#!/usr/bin/env python3
"""
Módulo Especialista: Coupa Enterprise (Trabalhador do Maestro)
Autor: Allan Simão
"""
import os
import re
import time
import json
import logging
import shutil
import tempfile
import sys
import threading
import random
import traceback
from datetime import datetime
from typing import List, Dict, Set, Optional, Any
from google import genai
from groq import Groq

import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Side, PatternFill

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.options import Options
from selenium.common.exceptions import NoAlertPresentException, NoSuchElementException
from selenium.webdriver.common.action_chains import ActionChains

from docx import Document
from docx.shared import Pt, RGBColor, Cm

import base64

logger = logging.getLogger("CoupaApp")
logger.setLevel(logging.INFO)
formatter = logging.Formatter('%(asctime)s [%(levelname)s] %(message)s', datefmt='%H:%M:%S')

if not logger.handlers:
    console_handler = logging.StreamHandler(sys.stdout)
    console_handler.setFormatter(formatter)
    logger.addHandler(console_handler)


class VerificacaoMixin:
    def fase_verificacao(self):
        logger.info("--- FASE 2: VERIFICAÇãO ---")
        try:
            self.garantir_logado(self.EVENTS_URL)
            self.driver.get(self.EVENTS_URL)
            time.sleep(3)
            
            # O get() acima já garante que estamos na primeira página.
            
            logger.info(f" Iniciando varredura total das últimas 40 páginas...")
            curr = 1
            limite_pag = 40 # Garante um limite mesmo se não estiver no __init__
            
            while True:
                if self.solicitacao_parada or curr > limite_pag: break
                logger.info(f"    Verificando Página {curr}...")
                self.sio.emit('relatar_progresso_coupa', {'mensagem': f"🔍 Verificando Página {curr}..."})
                
                try: 
                    WebDriverWait(self.driver, 20).until(EC.presence_of_element_located((By.TAG_NAME, "table")))
                except: 
                    break
                
                rows = self.driver.find_elements(By.CSS_SELECTOR, "table tbody tr")
                if not rows: break
                
                matches_pagina = 0
                
                for r in rows:
                    if self.solicitacao_parada: break
                    try:
                        num = r.find_element(By.CSS_SELECTOR, "a span.dt_open_link").text.strip()
                        matches_pagina += 1
                        qtd = 0
                        try: 
                            qtd = int(r.find_element(By.CLASS_NAME, "s-datatable-cell-num_responses").text)
                        except: pass
                        
                        if qtd > 0:
                            self.atualizar_status(num, True)
                            msg_ui = f"✅ {num} - RESPONDIDO"
                            self.sio.emit('relatar_progresso_coupa', {'mensagem': msg_ui})
                            logger.info(f"      {msg_ui}")
                            
                            if not self.db_eventos.get(num, {}).get('respondido'):
                                self.atualizar_status_respondido_jsonl(num)
                        else:
                            self.atualizar_status(num, False)
                            msg_ui = f"⏳ {num} - NÃO RESPONDIDO"
                            self.sio.emit('relatar_progresso_coupa', {'mensagem': msg_ui})
                            logger.info(f"      {msg_ui}")
                    except: 
                        continue
                
                logger.info(f"    Fim da Página {curr}. Matches encontrados: {matches_pagina}.")

                try:
                    next_btn = None
                    seletores = [
                        (By.CSS_SELECTOR, "a.next_page"),
                        (By.LINK_TEXT, "Próximo"),
                        (By.LINK_TEXT, "Next"),
                        (By.XPATH, "//a[contains(@class, 'next_page')]")
                    ]
                    
                    for metodo, query in seletores:
                        try:
                            el = self.driver.find_element(metodo, query)
                            if el.is_displayed():
                                next_btn = el; break
                        except: continue
                    
                    if not next_btn:
                        logger.info("    Fim da paginação (Botão não encontrado).")
                        break

                    cls = next_btn.get_attribute("class") or ""
                    try: p_cls = next_btn.find_element(By.XPATH, "..").get_attribute("class") or ""
                    except: p_cls = ""
                    
                    if "disabled" in cls or "disabled" in p_cls:
                        logger.info("    Fim da paginação (Botão desabilitado).")
                        break

                    self.driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", next_btn)
                    time.sleep(1)
                    self.driver.execute_script("arguments[0].click();", next_btn)
                    time.sleep(4) 
                    curr += 1
                except Exception as e:
                    logger.error(f" Erro crítico na paginação da Fase 2: {e}")
                    break
            
            logger.info("✅ Fase de verificação finalizada.")
            self.sio.emit('tarefa_concluida', {'evento': 'Verificação de Status', 'sucesso': True})
            
        except Exception as e:
            logger.error(f" Erro geral na Fase de Verificação: {e}")
            self.sio.emit('tarefa_concluida', {'evento': 'Verificação de Status', 'sucesso': False, 'erro': str(e)})

    def _get_vencidos(self):
        logger.info(" Lendo Excel para identificar eventos vencidos (Pente Fino)...")
        v = set()
        if not os.path.exists(self.ARQUIVO_PLANILHA_CONTROLE): 
            logger.error(" Arquivo de controle não encontrado para leitura.")
            return v
        
        wb = None # Inicia a variável vazia para evitar erros
        try:
            wb = load_workbook(self.ARQUIVO_PLANILHA_CONTROLE, data_only=True)
            ws = wb[self.NOME_ABA_ALVO] if self.NOME_ABA_ALVO in wb.sheetnames else wb.active
            
            logger.info(f"    O Excel reporta {ws.max_row} linhas preenchidas.")

            agora = datetime.now()
            hoje = agora.date()
            passou_do_horario = (agora.hour > 14) or (agora.hour == 14 and agora.minute >= 30)
            
            count_total = 0
            count_erro_data = 0
            
            for r in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=2):
                try:
                    if not r[0].value: continue 
                    val = r[1].value
                    dt_obj = None
                    
                    if isinstance(val, datetime):
                        dt_obj = val.date()
                    elif isinstance(val, str):
                        val = val.strip()
                        formatos = ["%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"]
                        for fmt in formatos:
                            try:
                                dt_obj = datetime.strptime(val.split(' ')[0], fmt.split(' ')[0]).date()
                                break
                            except: pass
                    
                    if dt_obj:
                        if dt_obj < hoje:
                            v.add(str(r[0].value).strip())
                            count_total += 1
                        elif dt_obj == hoje and passou_do_horario:
                            v.add(str(r[0].value).strip())
                            count_total += 1
                    else:
                        if val: count_erro_data += 1

                except: pass
            
            if count_erro_data > 0:
                logger.warning(f"    {count_erro_data} linhas ignoradas por data inválida ou formato desconhecido.")
            
            logger.info(f"✅ Total de eventos confirmados como VENCIDOS no Excel: {count_total}")
            
            wb.close()
            
            # <--- TRAVA DE SEGURANÇA 1: Libera a planilha na memória!
            return v
            
        except Exception as e:
            if 'wb' in locals() and wb:
                try: wb.close()
                except: pass 
            logger.error(f" Erro ao ler datas do Excel: {e}")
            if wb: 
                try: wb.close() # <--- TRAVA DE SEGURANÇA 2: Libera mesmo se der erro crítico!
                except: pass
            return v


